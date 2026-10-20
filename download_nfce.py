#!/usr/bin/env python3
"""Automação de download de NFC-e do portal Tec Fiscal / Linx.

Substitui a planilha VBA `Autom Download NF V4.xlsm`. Ver o design em
`docs/superpowers/specs/2026-08-10-download-nfce-design.md`.

Comandos:
    importar     lê as chaves da aba Base de um .xlsm e gera chaves.csv
    baixar       percorre a fila e baixa os PDFs do portal
    conferir     reconcilia o controle com os arquivos já existentes em disco
    status       mostra a contagem por status
    inspecionar  abre o navegador logado com o Playwright Inspector
"""

from __future__ import annotations

# ---------------------------------------------------------------------------
# 1. Imports e constantes
# ---------------------------------------------------------------------------

import argparse
import csv
import os
import re
import shutil
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

URL_LOGIN = "https://tecfiscal.linx.com.br/#/login"
URL_NOTAS = "https://tecfiscal.linx.com.br/app/emissao/notas"

TIMEOUT_PADRAO = 15          # segundos, sobrescrito por --timeout
TIMEOUT_LOGIN = 20           # segundos de espera pelo marcador de sessão ativa
ESPERA_LOGIN_MANUAL = 240    # segundos concedidos ao login manual (captcha inclusive)
TIMEOUT_DOWNLOAD = 60        # segundos de espera pelo evento de download
TIMEOUT_MODAL = 3            # segundos de espera pelo modal pós-download
SONDAGEM_MS = 500            # fatia de espera por localizador candidato
PAUSA_ENTRE_TENTATIVAS = 3   # segundos entre tentativas da mesma chave
TENTATIVAS_ANTES_RECARREGAR = 3  # erros seguidos até recarregar a página (bug de layout)

# Navegadores tentados em ordem. None = Chromium baixado pelo Playwright.
CANAIS_NAVEGADOR = (None, "msedge", "chrome")

LOTE_PADRAO = 150
TENTATIVAS_PADRAO = 6

CHAVES_PADRAO = "chaves.csv"
CONTROLE_PADRAO = "controle.csv"
DESTINO_PADRAO = "./pdfs"
DIR_SCREENSHOTS = Path("logs/erros")
LOG_ERROS = Path("logs/erros.log")

STATUS_PENDENTE = "pendente"
STATUS_BAIXADO = "baixado"
STATUS_NAO_ENCONTRADO = "nao_encontrado"
STATUS_ERRO = "erro"
ORDEM_STATUS = (STATUS_PENDENTE, STATUS_BAIXADO, STATUS_NAO_ENCONTRADO, STATUS_ERRO)

COLUNAS_CONTROLE = ("chave", "status", "tentativas", "arquivo", "erro", "atualizado_em")
CHAVE_RE = re.compile(r"\d{44}")
LIMITE_MENSAGEM_ERRO = 200
ENCODING = "utf-8-sig"
# Backoff do replace do controle. Ver _substituir_com_retry.
ESPERAS_REGRAVACAO = (0.0, 0.1, 0.3, 0.6, 1.0)


# ---------------------------------------------------------------------------
# 2. SELETORES
# ---------------------------------------------------------------------------
#
# Cada ação mapeia para uma LISTA ORDENADA de localizadores; o código usa o
# primeiro que resolver para um elemento visível. Ordem de preferência:
#   papel/acessibilidade  ->  CSS estável (nz-, lx-, data-testid)  ->  XPath
# posicional da especificação (frágil, último recurso).
#
# Prefixos aceitos: "papel:<role>:<nome>", "rotulo:", "placeholder:",
# "texto:", "titulo:". Sem prefixo, a string vai direta para page.locator()
# (aceita CSS e "xpath=...").
#
# Esta é a parte com maior probabilidade de quebrar: se o layout do portal
# mudar, conserta-se aqui. Use `python download_nfce.py inspecionar` para
# descobrir os localizadores novos.

SELETORES: dict[str, list[str]] = {
    # Não use "Filtrar" aqui: no portal esse é o botão de APLICAR, e ele venceria
    # a lista antes do botão que abre o painel.
    "botao_abrir_filtro": [
        "papel:button:Filtro",
        "lx-page-filter button[title*='iltro' i]",
        "lx-page-filter > div > div:nth-child(1) button",
        "xpath=//lx-page-filter/div/div[1]/button",
    ],
    "campo_chave": [
        "placeholder:Chave",
        "rotulo:Chave de acesso",
        "app-lx-filter-text input[type='text']",
        "lx-page-filter input[type='text']",
        "xpath=//lx-page-filter/div/div[2]/form/app-lx-filter-text[2]/div/div[2]/input",
    ],
    "botao_aplicar_filtro": [
        "papel:button:Pesquisar",
        "papel:button:Aplicar",
        "papel:button:Filtrar",
        "lx-page-filter button[type='submit']",
        "xpath=//lx-page-filter/div/div[1]/div/div/div[2]/button[2]",
    ],
    "spinner": [
        "nz-spin .ant-spin-spinning",
        ".ant-spin-spinning",
        "nz-spin[nzspinning='true']",
    ],
    # `tr:not(.ant-table-placeholder)` é essencial: em tabelas ng-zorro o estado
    # vazio também é uma <tr> dentro da tbody, e casá-la faria "não encontrado"
    # virar "encontrado".
    "linha_resultado": [
        "linx-table nz-table tbody tr[data-row-key]",
        "nz-table tbody tr[data-row-key]",
        "linx-table tbody tr:not(.ant-table-placeholder)",
        "nz-table tbody tr:not(.ant-table-placeholder)",
        "xpath=//linx-table//tbody/tr[not(contains(@class,'placeholder'))][1]",
    ],
    "resultado_vazio": [
        "nz-embed-empty",
        ".ant-empty",
        "texto:Nenhum registro",
        "texto:Nenhum resultado",
        "texto:Sem dados",
    ],
    "checkbox_nota": [
        "linx-table nz-table tbody tr[data-row-key] input[type='checkbox']",
        "nz-table tbody tr[data-row-key] input[type='checkbox']",
        "nz-table tbody tr:not(.ant-table-placeholder) input[type='checkbox']",
        "xpath=//linx-table//tbody/tr[not(contains(@class,'placeholder'))][1]/td[1]//input[@type='checkbox']",
        "nz-table tbody tr[data-row-key] label.ant-checkbox-wrapper",
        "nz-table tbody tr:not(.ant-table-placeholder) td:first-child label",
    ],
    # A barra de seleção tem QUATRO botões de download: XML, DANFE, Resumo Sem
    # Valor Fiscal e Relatório de Notas. O nome acessível de todos começa com a
    # ligadura "download" do ícone, então casar por "Download" pega o primeiro em
    # ordem de DOM (o XML). Cada formato precisa casar pelo próprio rótulo.
    "botao_download": [
        "papel:button:DANFE",
        "xpath=//linx-table-toolbar//button[contains(normalize-space(.),'DANFE')]",
    ],
    "botao_download_xml": [
        "papel:button:XML",
        "xpath=//linx-table-toolbar//button[contains(normalize-space(.),'XML')]",
    ],
    # Toast que o portal mostra quando não consegue gerar o documento (ex.: "Não
    # foi possível gerar a danfe... O valor da coluna 'xCpl' na tabela 'entrega'
    # é DBNull"). Nesse caso o download nunca começa.
    "notificacao_erro": [
        ".ant-notification-notice:has(.ant-notification-notice-icon-error)",
        ".ant-notification-notice-error",
        ".ant-message-error",
        "[role='alert']:has-text('Não foi possível')",
        "texto:Não foi possível gerar",
    ],
    "fechar_notificacao": [
        ".ant-notification-notice-close",
        ".ant-notification-notice button[aria-label*='echar' i]",
    ],
    "botao_fechar_modal": [
        "papel:button:Fechar",
        ".ant-modal-close",
        "nz-modal button.ant-modal-close",
        "xpath=/html/body/div/div[2]/div/div/div[2]/div/div/div[1]/div/button",
    ],
    "marcador_logado": [
        "lx-page-filter",
        "linx-table",
        "papel:button:Filtro",
    ],
    # Confirmados contra o DOM real em 2026-08-10: o login é por e-mail.
    "campo_usuario": [
        "#email",
        "input[formcontrolname='Email']",
        "placeholder:Digite seu e-mail de acesso",
        "rotulo:E-mail",
        "input[type='text']",
    ],
    "campo_senha": [
        "#arialabelpassword",
        "rotulo:Senha",
        "placeholder:Senha",
        "input[type='password']",
    ],
    "botao_entrar": [
        "papel:button:Entrar",
        "papel:button:Acessar",
        "papel:button:Login",
        "form button[type='submit']",
    ],
    # Banner de LGPD no rodapé. Fica sobre a página e pode interceptar cliques.
    "botao_cookies": [
        "papel:button:Aceitar cookies",
        "papel:button:Aceitar",
    ],
    "seletor_filial": [
        "papel:combobox:Filial",
        "nz-select[formcontrolname='filial']",
        "lx-page-header nz-select",
        "nz-select",
    ],
}


class ErroPortal(RuntimeError):
    """Falha técnica durante a interação com o portal."""


class SeletorNaoEncontrado(ErroPortal):
    """Nenhum localizador da lista resolveu para um elemento visível."""


class ErroGeracaoPortal(ErroPortal):
    """O portal recusou gerar o documento da nota (erro nos dados dela).

    É determinístico: repetir na mesma rodada só gasta tempo. A chave fica como
    `erro` e volta para a fila na próxima execução.
    """


# ---------------------------------------------------------------------------
# 3. Configuração (.env)
# ---------------------------------------------------------------------------


@dataclass
class Config:
    usuario: str = ""
    senha: str = ""
    cnpj: str = ""

    @property
    def tem_credenciais(self) -> bool:
        return bool(self.usuario and self.senha)


def _carregar_env_manual(caminho: Path) -> None:
    """Fallback mínimo para quando python-dotenv não estiver instalado."""
    if not caminho.exists():
        return
    for linha in caminho.read_text(encoding=ENCODING).splitlines():
        linha = linha.strip()
        if not linha or linha.startswith("#") or "=" not in linha:
            continue
        nome, _, valor = linha.partition("=")
        os.environ.setdefault(nome.strip(), valor.strip().strip("'\""))


def carregar_config(caminho: str | Path = ".env") -> Config:
    caminho = Path(caminho)
    try:
        from dotenv import load_dotenv

        load_dotenv(caminho)
    except ImportError:
        _carregar_env_manual(caminho)
    return Config(
        usuario=os.environ.get("TECFISCAL_USUARIO", "").strip(),
        senha=os.environ.get("TECFISCAL_SENHA", ""),
        cnpj=os.environ.get("TECFISCAL_CNPJ", "").strip(),
    )


# ---------------------------------------------------------------------------
# 4. Funções de CSV (chaves e controle)
# ---------------------------------------------------------------------------


@dataclass
class Registro:
    chave: str
    status: str = STATUS_PENDENTE
    tentativas: int = 0
    arquivo: str = ""
    erro: str = ""
    atualizado_em: str = ""

    def marcar(self, status: str, *, arquivo: str = "", erro: str = "", tentativas: int = 0) -> None:
        self.status = status
        self.tentativas += tentativas
        self.arquivo = arquivo
        self.erro = truncar(erro)
        self.atualizado_em = agora_iso()


def agora_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def truncar(mensagem: str, limite: int = LIMITE_MENSAGEM_ERRO) -> str:
    """Achata quebras de linha e corta em `limite` caracteres."""
    texto = " ".join(str(mensagem).split())
    return texto[:limite]


def normalizar_chave(valor: object) -> str:
    """Remove máscara e devolve só os dígitos. Números viram inteiro sem notação."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"\D", "", str(valor))


def ler_chaves(caminho: str | Path) -> list[str]:
    """Lê chaves.csv (coluna `chave`, com cabeçalho), sem duplicatas."""
    caminho = Path(caminho)
    if not caminho.exists():
        raise SystemExit(f"Arquivo de chaves não encontrado: {caminho}")
    chaves: list[str] = []
    vistas: set[str] = set()
    with caminho.open(encoding=ENCODING, newline="") as f:
        for linha in csv.DictReader(f):
            chave = normalizar_chave(linha.get("chave"))
            if len(chave) == 44 and chave not in vistas:
                vistas.add(chave)
                chaves.append(chave)
    return chaves


def gravar_chaves(caminho: str | Path, chaves: list[str]) -> None:
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with caminho.open("w", encoding=ENCODING, newline="") as f:
        escritor = csv.writer(f)
        escritor.writerow(["chave"])
        escritor.writerows([[c] for c in chaves])


def carregar_controle(caminho: str | Path) -> dict[str, Registro]:
    """Lê controle.csv. Arquivo ausente vira controle vazio."""
    caminho = Path(caminho)
    controle: dict[str, Registro] = {}
    if not caminho.exists():
        return controle
    with caminho.open(encoding=ENCODING, newline="") as f:
        for linha in csv.DictReader(f):
            chave = normalizar_chave(linha.get("chave"))
            if not chave:
                continue
            try:
                tentativas = int(linha.get("tentativas") or 0)
            except ValueError:
                tentativas = 0
            controle[chave] = Registro(
                chave=chave,
                status=(linha.get("status") or STATUS_PENDENTE).strip() or STATUS_PENDENTE,
                tentativas=tentativas,
                arquivo=(linha.get("arquivo") or "").strip(),
                erro=(linha.get("erro") or "").strip(),
                atualizado_em=(linha.get("atualizado_em") or "").strip(),
            )
    return controle


def _substituir_com_retry(temporario: Path, caminho: Path) -> None:
    """`os.replace` com retry e degradação para cópia não-atômica.

    O controle é regravado a cada chave — milhares de vezes por rodada. Em pasta
    sincronizada (OneDrive) ou com antivírus ativo, o destino fica bloqueado por
    instantes e o replace estoura `PermissionError: Acesso negado`. Insistir
    resolve; se nem assim, sobrescrever no lugar é melhor que perder o progresso.
    """
    for espera in ESPERAS_REGRAVACAO:
        try:
            os.replace(temporario, caminho)
            return
        except PermissionError:
            time.sleep(espera)
    shutil.copyfile(temporario, caminho)
    temporario.unlink(missing_ok=True)


def gravar_controle(caminho: str | Path, controle: dict[str, Registro]) -> None:
    """Grava o controle inteiro de forma atômica (temporário + replace)."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    temporario = caminho.with_suffix(caminho.suffix + ".tmp")
    with temporario.open("w", encoding=ENCODING, newline="") as f:
        escritor = csv.writer(f)
        escritor.writerow(COLUNAS_CONTROLE)
        for reg in controle.values():
            escritor.writerow(
                [reg.chave, reg.status, reg.tentativas, reg.arquivo, reg.erro, reg.atualizado_em]
            )
    _substituir_com_retry(temporario, caminho)


def sincronizar_controle(controle: dict[str, Registro], chaves: list[str]) -> int:
    """Insere como `pendente` as chaves da entrada ausentes do controle.

    Chaves que só existem no controle são preservadas — não se apaga histórico.
    Devolve quantas foram inseridas.
    """
    novas = 0
    for chave in chaves:
        if chave not in controle:
            controle[chave] = Registro(chave=chave, status=STATUS_PENDENTE, atualizado_em=agora_iso())
            novas += 1
    return novas


def montar_fila(controle: dict[str, Registro], chaves: list[str], retentar: bool = False) -> list[str]:
    """Fila da rodada: `pendente` + `erro`, mais `nao_encontrado` se `retentar`."""
    elegiveis = {STATUS_PENDENTE, STATUS_ERRO}
    if retentar:
        elegiveis.add(STATUS_NAO_ENCONTRADO)
    fila = []
    for chave in chaves:
        reg = controle.get(chave)
        if reg is None or reg.status in elegiveis:
            fila.append(chave)
    return fila


def contar_status(controle: dict[str, Registro]) -> dict[str, int]:
    contagem = {status: 0 for status in ORDEM_STATUS}
    for reg in controle.values():
        contagem[reg.status] = contagem.get(reg.status, 0) + 1
    return contagem


def imprimir_resumo(controle: dict[str, Registro], titulo: str = "Resumo") -> None:
    contagem = contar_status(controle)
    print(f"\n{titulo}")
    print("-" * 32)
    for status in ORDEM_STATUS:
        print(f"  {status:<16} {contagem.get(status, 0):>6}")
    for status, total in contagem.items():
        if status not in ORDEM_STATUS:
            print(f"  {status:<16} {total:>6}")
    print("-" * 32)
    print(f"  {'total':<16} {len(controle):>6}")


# ---------------------------------------------------------------------------
# 5. Varredura de diretório
# ---------------------------------------------------------------------------


def extrair_chave_do_nome(nome: str) -> str | None:
    """Primeira sequência de 44 dígitos no nome do arquivo, ou None."""
    achado = CHAVE_RE.search(nome)
    return achado.group(0) if achado else None


def varrer_destino(destino: str | Path) -> tuple[dict[str, Path], list[Path]]:
    """Varre `destino` recursivamente.

    Devolve (mapa chave -> arquivo, arquivos cujo nome não contém chave).
    Quando a mesma chave aparece mais de uma vez, o primeiro achado vence.
    """
    destino = Path(destino)
    mapa: dict[str, Path] = {}
    sem_chave: list[Path] = []
    if not destino.exists():
        return mapa, sem_chave
    for arquivo in sorted(destino.rglob("*")):
        if not arquivo.is_file():
            continue
        chave = extrair_chave_do_nome(arquivo.name)
        if chave is None:
            sem_chave.append(arquivo)
        elif chave not in mapa:
            mapa[chave] = arquivo
    return mapa, sem_chave


# ---------------------------------------------------------------------------
# 6. Importação do .xlsm
# ---------------------------------------------------------------------------


@dataclass
class ResultadoImportacao:
    chaves: list[str] = field(default_factory=list)
    total_lido: int = 0
    invalidas: list[tuple[int, str]] = field(default_factory=list)
    duplicadas: int = 0


def importar_chaves_planilha(caminho: str | Path, aba: str = "Base") -> ResultadoImportacao:
    """Extrai as chaves da coluna A, da linha 2 em diante, da aba indicada."""
    from openpyxl import load_workbook

    caminho = Path(caminho)
    if not caminho.exists():
        raise SystemExit(f"Planilha não encontrada: {caminho}")

    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        if aba not in livro.sheetnames:
            raise SystemExit(
                f"Aba '{aba}' não existe em {caminho.name}. Abas: {', '.join(livro.sheetnames)}"
            )
        planilha = livro[aba]
        resultado = ResultadoImportacao()
        vistas: set[str] = set()
        for numero, (celula,) in enumerate(
            planilha.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True), start=2
        ):
            if celula is None or str(celula).strip() == "":
                continue
            resultado.total_lido += 1
            chave = normalizar_chave(celula)
            if len(chave) != 44:
                resultado.invalidas.append((numero, truncar(celula, 60)))
            elif chave in vistas:
                resultado.duplicadas += 1
            else:
                vistas.add(chave)
                resultado.chaves.append(chave)
        return resultado
    finally:
        livro.close()


# ---------------------------------------------------------------------------
# 7. Funções de portal
# ---------------------------------------------------------------------------


def resolver(page, descritor: str):
    """Traduz um descritor de SELETORES em um Locator do Playwright."""
    prefixo, _, resto = descritor.partition(":")
    if prefixo == "papel":
        papel, _, nome = resto.partition(":")
        return page.get_by_role(papel, name=nome)
    if prefixo == "rotulo":
        return page.get_by_label(resto)
    if prefixo == "placeholder":
        return page.get_by_placeholder(resto)
    if prefixo == "texto":
        return page.get_by_text(resto)
    if prefixo == "titulo":
        return page.get_by_title(resto)
    return page.locator(descritor)


def localizar(page, acao: str, timeout_ms: int):
    """Primeiro localizador visível da lista de `acao`.

    Sonda os candidatos em rodadas curtas até esgotar o timeout, de modo que
    qualquer um deles possa aproveitar a janela inteira de espera.
    """
    descritores = SELETORES[acao]
    limite = time.monotonic() + timeout_ms / 1000
    while True:
        for descritor in descritores:
            try:
                locator = resolver(page, descritor).first
                locator.wait_for(state="visible", timeout=SONDAGEM_MS)
                return locator
            except Exception:
                continue
        if time.monotonic() >= limite:
            raise SeletorNaoEncontrado(
                f"Não encontrei '{acao}' na página. Rode `inspecionar` e ajuste SELETORES."
            )


def visivel(page, acao: str, timeout_ms: int = SONDAGEM_MS) -> bool:
    for descritor in SELETORES[acao]:
        try:
            resolver(page, descritor).first.wait_for(state="visible", timeout=timeout_ms)
            return True
        except Exception:
            continue
    return False


def clicar(page, acao: str, timeout_ms: int) -> None:
    localizar(page, acao, timeout_ms).click(timeout=timeout_ms)


def preencher(page, acao: str, texto: str, timeout_ms: int) -> None:
    campo = localizar(page, acao, timeout_ms)
    campo.fill("", timeout=timeout_ms)
    campo.fill(texto, timeout=timeout_ms)


def aguardar_spinner(page, timeout_ms: int) -> None:
    """Espera o indicador de carregamento sumir. Ausência dele é sucesso."""
    limite = time.monotonic() + timeout_ms / 1000
    for descritor in SELETORES["spinner"]:
        while time.monotonic() < limite:
            try:
                locator = resolver(page, descritor).first
                if locator.count() == 0 or not locator.is_visible():
                    break
            except Exception:
                break
            page.wait_for_timeout(200)


def aguardar_resultado(page, timeout_ms: int) -> bool:
    """True se a busca trouxe linha, False se trouxe o indicador de vazio.

    "Tabela vazia" é um resultado legítimo e não pode ser confundido com
    "tabela ainda carregando" — por isso o spinner é esperado antes, e depois
    corre-se linha x vazio, decidindo pelo primeiro que aparecer.
    """
    aguardar_spinner(page, timeout_ms)
    limite = time.monotonic() + timeout_ms / 1000
    while True:
        if visivel(page, "linha_resultado", 300):
            return True
        if visivel(page, "resultado_vazio", 300):
            # Confirma: o indicador de vazio pode aparecer por um instante antes
            # de a linha chegar. Só se nenhuma linha surgir é "não encontrado".
            page.wait_for_timeout(500)
            return visivel(page, "linha_resultado", 300)
        if time.monotonic() >= limite:
            raise ErroPortal("O resultado da busca não estabilizou (nem linha, nem vazio).")


def fechar_modal(page) -> None:
    """Fecha o modal residual do pós-download. Ausência não é erro."""
    try:
        localizar(page, "botao_fechar_modal", TIMEOUT_MODAL * 1000).click(timeout=TIMEOUT_MODAL * 1000)
    except Exception:
        pass


def primeiro_visivel_agora(page, acao: str):
    """Primeiro localizador de `acao` visível neste instante, sem esperar."""
    for descritor in SELETORES[acao]:
        try:
            locator = resolver(page, descritor).first
            if locator.count() and locator.is_visible():
                return locator
        except Exception:
            continue
    return None


def mensagem_erro_portal(page) -> str:
    """Texto do toast de erro do portal, ou "" se não houver nenhum."""
    locator = primeiro_visivel_agora(page, "notificacao_erro")
    if locator is None:
        return ""
    try:
        return " ".join(locator.inner_text(timeout=1000).split()) or "Erro sem mensagem"
    except Exception:
        return "Erro sem mensagem"


def fechar_notificacoes(page) -> None:
    """Fecha os toasts abertos, para não confundirem a próxima nota."""
    for descritor in SELETORES["fechar_notificacao"]:
        try:
            botoes = resolver(page, descritor)
            for i in range(botoes.count()):
                botao = botoes.nth(i)
                if botao.is_visible():
                    botao.click(timeout=2000)
        except Exception:
            continue


def baixar_arquivo(
    page, chave: str, destino: Path, timeout_ms: int, acao: str = "botao_download"
) -> Path:
    """Clica no download e corre o evento de download x toast de erro do portal.

    Quando o portal não consegue gerar o documento ele não dispara download:
    mostra um toast de erro. Sem olhar para o toast, a espera iria até o
    timeout do download em toda tentativa.
    """
    fechar_notificacoes(page)
    downloads: list = []

    # Precisa ser função Python: o Playwright grava um atributo no ouvinte, e
    # método embutido (como `downloads.append`) não aceita atributo.
    def ao_baixar(download) -> None:
        downloads.append(download)

    page.on("download", ao_baixar)
    try:
        clicar(page, acao, timeout_ms)
        limite = time.monotonic() + TIMEOUT_DOWNLOAD
        while not downloads:
            mensagem = mensagem_erro_portal(page)
            if mensagem:
                raise ErroGeracaoPortal(mensagem)
            if time.monotonic() >= limite:
                raise ErroPortal(f"O download não começou em {TIMEOUT_DOWNLOAD}s.")
            # wait_for_timeout (e não time.sleep) para o Playwright processar o evento.
            page.wait_for_timeout(300)
    finally:
        page.remove_listener("download", ao_baixar)
    download = downloads[0]
    extensao = Path(download.suggested_filename).suffix or ".pdf"
    destino_arquivo = destino / f"{chave}{extensao}"
    download.save_as(str(destino_arquivo))
    if not destino_arquivo.exists() or destino_arquivo.stat().st_size == 0:
        raise ErroPortal(f"Download vazio ou não salvo: {destino_arquivo.name}")
    return destino_arquivo


def na_tela_de_login(page) -> bool:
    """Há campo de senha visível, ou a URL ainda é a de login."""
    if "/login" in (page.url or ""):
        return True
    try:
        campo = page.locator("input[type='password']").first
        return campo.count() > 0 and campo.is_visible()
    except Exception:
        return False


def sessao_ativa(page, timeout_ms: int) -> bool:
    """Marcador de sessão: estar fora do login e na tela de notas."""
    limite = time.monotonic() + timeout_ms / 1000
    while True:
        if not na_tela_de_login(page):
            if visivel(page, "marcador_logado", 300):
                return True
            if "emissao/notas" in (page.url or ""):
                return True
        if time.monotonic() >= limite:
            return False
        page.wait_for_timeout(300)


def confirmar_sessao(page, timeout_ms: int) -> bool:
    """Confirma a sessão, indo para a tela de notas se preciso.

    Depois do login o portal cai na tela inicial, não na de notas — e é lá que
    vivem os marcadores. Sair do login não basta: tem que navegar.
    """
    if sessao_ativa(page, 2000):
        return True
    if na_tela_de_login(page):
        return False
    try:
        page.goto(URL_NOTAS, wait_until="domcontentloaded")
    except Exception:
        return False
    aceitar_cookies(page)
    return sessao_ativa(page, timeout_ms)


def selecionar_filial(page, cnpj: str, timeout_ms: int) -> None:
    """Seleção opcional de filial por CNPJ. Falha aqui não derruba a rodada."""
    try:
        clicar(page, "seletor_filial", timeout_ms)
        digitos = re.sub(r"\D", "", cnpj)
        for alvo in (cnpj, digitos):
            opcao = page.get_by_text(alvo, exact=False).first
            try:
                opcao.wait_for(state="visible", timeout=2000)
                opcao.click(timeout=timeout_ms)
                return
            except Exception:
                continue
        print(f"  aviso: filial {cnpj} não encontrada no seletor; seguindo com a filial atual.")
    except Exception as exc:
        print(f"  aviso: não consegui selecionar a filial ({truncar(exc, 100)}); seguindo.")


def aceitar_cookies(page) -> None:
    """Dispensa o banner de LGPD, que sobrepõe a página. Ausência não é erro."""
    try:
        localizar(page, "botao_cookies", 3000).click(timeout=3000)
    except Exception:
        pass


def esperar_login_manual(page, segundos: int = ESPERA_LOGIN_MANUAL) -> bool:
    """Aguarda o usuário logar na janela do navegador.

    Divergência deliberada em relação ao design (§5.2), que previa `input()` no
    modelo do MsgBox do VBA: aqui a sessão é detectada por sondagem. Além de
    dispensar um ENTER, isso faz o modo manual funcionar mesmo quando a entrada
    padrão não é um terminal (script agendado, chamada por outra ferramenta).
    """
    print(
        f"Faça login na janela do navegador (inclusive o captcha).\n"
        f"Aguardando até {segundos}s — sigo sozinho assim que a sessão aparecer.",
        flush=True,
    )
    limite = time.monotonic() + segundos
    while time.monotonic() < limite:
        # confirmar_sessao navega para a tela de notas assim que o login sai do
        # caminho — o portal não cai lá sozinho depois de autenticar.
        if confirmar_sessao(page, 8000):
            print("Sessão detectada. Seguindo.", flush=True)
            return True
        onde = "tela de login" if na_tela_de_login(page) else page.url
        print(f"  ... aguardando ({int(limite - time.monotonic())}s) — {onde}", flush=True)
        time.sleep(3)
    return False


def fazer_login(
    page, cfg: Config, headless: bool, timeout_ms: int, espera_login: int = ESPERA_LOGIN_MANUAL
) -> None:
    """Login híbrido: tenta automático com o .env, cai para manual se falhar.

    Atenção: em 2026-08-10 a tela de login do portal exibia reCAPTCHA. Com o
    captcha ativo o preenchimento automático não conclui o login sozinho — o
    caminho real é o manual, e por isso vale usar `--perfil` para não repetir
    o captcha a cada reciclagem de sessão.
    """
    # Com --perfil os cookies sobrevivem entre execuções: se a sessão ainda vale,
    # não há login nenhum a fazer. Sem perfil, isso só redireciona para o login.
    page.goto(URL_NOTAS, wait_until="domcontentloaded")
    aceitar_cookies(page)
    if confirmar_sessao(page, 8000):
        print("Sessão reaproveitada do perfil — sem login.")
        return

    page.goto(URL_LOGIN, wait_until="domcontentloaded")
    aceitar_cookies(page)

    if cfg.tem_credenciais:
        try:
            preencher(page, "campo_usuario", cfg.usuario, timeout_ms)
            preencher(page, "campo_senha", cfg.senha, timeout_ms)
            clicar(page, "botao_entrar", timeout_ms)
        except Exception:
            # Pode já estar logado, ou a tela de login mudou: o marcador decide.
            pass
    elif not headless:
        print("Sem credenciais no .env — o login será manual.")

    if confirmar_sessao(page, TIMEOUT_LOGIN * 1000):
        return

    if headless:
        raise SystemExit(
            "Login automático falhou e --headless não permite intervenção.\n"
            "Rode sem --headless, ou confira TECFISCAL_USUARIO/TECFISCAL_SENHA no .env."
        )

    if not esperar_login_manual(page, espera_login):
        raise SystemExit(f"Sessão não confirmada em {espera_login}s de login manual. Abortando.")


class Sessao:
    """Ciclo de vida do navegador. Reciclada a cada --lote chaves."""

    def __init__(
        self,
        cfg: Config,
        destino: Path,
        headless: bool,
        timeout_ms: int,
        perfil: str | None = None,
        espera_login: int = ESPERA_LOGIN_MANUAL,
    ):
        self.cfg = cfg
        self.destino = destino
        self.headless = headless
        self.timeout_ms = timeout_ms
        self.espera_login = espera_login
        # Diretório de perfil do Chromium. Com ele os cookies sobrevivem à
        # reciclagem de sessão, e o captcha do login não se repete a cada lote.
        self.perfil = perfil
        self._playwright = None
        self._navegador = None
        self._contexto = None
        self.page = None

    def abrir(self) -> None:
        from playwright.sync_api import sync_playwright

        self._playwright = sync_playwright().start()
        self._lancar_navegador()
        self._contexto.set_default_timeout(self.timeout_ms)
        fazer_login(self.page, self.cfg, self.headless, self.timeout_ms, self.espera_login)
        if self.cfg.cnpj:
            selecionar_filial(self.page, self.cfg.cnpj, self.timeout_ms)

    def _lancar_navegador(self) -> None:
        """Chromium do Playwright; sem ele, cai para o Edge e depois o Chrome.

        Em rede corporativa o download do Chromium (`playwright install`) pode
        ser bloqueado pelo proxy. O Edge vem em todo Windows, então o robô roda
        mesmo assim.
        """
        primeiro_erro = None
        for canal in CANAIS_NAVEGADOR:
            opcoes = {"headless": self.headless}
            if canal:
                opcoes["channel"] = canal
            try:
                if self.perfil:
                    Path(self.perfil).mkdir(parents=True, exist_ok=True)
                    self._contexto = self._playwright.chromium.launch_persistent_context(
                        self.perfil, accept_downloads=True, **opcoes
                    )
                    paginas = self._contexto.pages
                    self.page = paginas[0] if paginas else self._contexto.new_page()
                else:
                    self._navegador = self._playwright.chromium.launch(**opcoes)
                    self._contexto = self._navegador.new_context(accept_downloads=True)
                    self.page = self._contexto.new_page()
            except Exception as exc:
                primeiro_erro = primeiro_erro or exc
                continue
            if canal:
                print(f"  aviso: Chromium do Playwright indisponível; usando o navegador '{canal}'.")
            return
        raise SystemExit(
            "Não consegui abrir nenhum navegador (Chromium, Edge ou Chrome).\n"
            f"Erro: {truncar(primeiro_erro, 300)}\n"
            "Rode `python -m playwright install chromium` ou instale o Microsoft Edge."
        )

    def fechar(self) -> None:
        for alvo in (self._contexto, self._navegador, self._playwright):
            if alvo is None:
                continue
            try:
                alvo.stop() if hasattr(alvo, "stop") else alvo.close()
            except Exception:
                pass
        self._contexto = self._navegador = self._playwright = self.page = None

    def reciclar(self) -> None:
        self.fechar()
        self.abrir()

    def __enter__(self) -> "Sessao":
        self.abrir()
        return self

    def __exit__(self, *_exc) -> None:
        self.fechar()


@dataclass
class ResultadoChave:
    status: str
    arquivo: str = ""
    erro: str = ""
    tentativas: int = 0


def processar_chave(
    page,
    chave: str,
    destino: Path,
    timeout_ms: int,
    tentativas: int,
    acao_download: str = "botao_download",
) -> ResultadoChave:
    """Passos 2 a 9 do laço, com retry. Nunca levanta exceção de portal."""
    ultimo_erro = ""
    for numero in range(1, tentativas + 1):
        try:
            clicar(page, "botao_abrir_filtro", timeout_ms)
            preencher(page, "campo_chave", chave, timeout_ms)
            clicar(page, "botao_aplicar_filtro", timeout_ms)
            if not aguardar_resultado(page, timeout_ms):
                # Resultado válido da busca, não é falha: não consome tentativas.
                return ResultadoChave(STATUS_NAO_ENCONTRADO, tentativas=0)
            clicar(page, "checkbox_nota", timeout_ms)
            arquivo = baixar_arquivo(page, chave, destino, timeout_ms, acao_download)
            fechar_modal(page)
            return ResultadoChave(STATUS_BAIXADO, arquivo=arquivo.name, tentativas=numero)
        except ErroGeracaoPortal as exc:
            # Erro nos dados da nota: não adianta insistir agora. Segue para a próxima.
            return ResultadoChave(STATUS_ERRO, erro=f"Portal: {exc}", tentativas=numero)
        except Exception as exc:
            ultimo_erro = f"{type(exc).__name__}: {exc}"
            if numero < tentativas:
                if numero % TENTATIVAS_ANTES_RECARREGAR == 0:
                    # Erro de layout pode deixar a página num estado que nunca se
                    # recupera sozinho; recarregar reinicia o processo e evita
                    # ficar preso repetindo o mesmo erro até esgotar as tentativas.
                    try:
                        page.goto(URL_NOTAS, wait_until="domcontentloaded", timeout=timeout_ms)
                    except Exception:
                        pass
                time.sleep(PAUSA_ENTRE_TENTATIVAS)
    return ResultadoChave(STATUS_ERRO, erro=ultimo_erro, tentativas=tentativas)


def salvar_screenshot(page, chave: str) -> None:
    try:
        DIR_SCREENSHOTS.mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(DIR_SCREENSHOTS / f"{chave}.png"), full_page=True)
    except Exception:
        pass


def registrar_erro(chave: str, mensagem: str, caminho: Path = LOG_ERROS) -> None:
    """Acrescenta o erro completo ao log. O controle guarda só um resumo truncado."""
    try:
        caminho.parent.mkdir(parents=True, exist_ok=True)
        with caminho.open("a", encoding="utf-8") as f:
            f.write(f"{agora_iso()}\t{chave}\t{' '.join(str(mensagem).split())}\n")
    except Exception:
        pass


# ---------------------------------------------------------------------------
# 8. Loop principal
# ---------------------------------------------------------------------------


def comando_baixar(args: argparse.Namespace) -> int:
    if args.lote < 1:
        raise SystemExit("--lote precisa ser pelo menos 1.")
    if args.tentativas < 1:
        raise SystemExit("--tentativas precisa ser pelo menos 1.")

    cfg = carregar_config()
    destino = Path(args.destino)
    destino.mkdir(parents=True, exist_ok=True)
    timeout_ms = int(args.timeout * 1000)
    acao_download = "botao_download" if args.formato == "danfe" else "botao_download_xml"

    chaves = ler_chaves(args.chaves)
    controle = carregar_controle(args.controle)
    novas = sincronizar_controle(controle, chaves)
    if novas:
        print(f"{novas} chave(s) nova(s) inserida(s) como pendente.")
        gravar_controle(args.controle, controle)

    fila = montar_fila(controle, chaves, args.retentar)
    if args.limite:
        fila = fila[: args.limite]
    if not fila:
        print("Nada a processar.")
        imprimir_resumo(controle)
        return 0

    em_disco, _ = varrer_destino(destino)
    print(f"{len(fila)} chave(s) na fila | {len(em_disco)} arquivo(s) já no destino.")

    processadas = 0
    erros_rodada = 0
    interrompido = False
    sessao = Sessao(cfg, destino, args.headless, timeout_ms, args.perfil, args.espera_login)
    try:
        # Dentro do try: se o login abortar, o `finally` ainda fecha o navegador.
        sessao.abrir()
        for posicao, chave in enumerate(fila, start=1):
            # Curto-circuito de idempotência: o arquivo já está em disco.
            if chave in em_disco:
                controle[chave].marcar(STATUS_BAIXADO, arquivo=em_disco[chave].name)
                gravar_controle(args.controle, controle)
                print(f"[{posicao}/{len(fila)}] {chave} já em disco")
                continue

            if processadas and processadas % args.lote == 0:
                print(f"-- reciclando a sessão após {processadas} chaves --")
                sessao.reciclar()

            resultado = processar_chave(
                sessao.page, chave, destino, timeout_ms, args.tentativas, acao_download
            )
            processadas += 1
            controle[chave].marcar(
                resultado.status,
                arquivo=resultado.arquivo,
                erro=resultado.erro,
                tentativas=resultado.tentativas,
            )
            gravar_controle(args.controle, controle)

            if resultado.status == STATUS_BAIXADO:
                em_disco[chave] = destino / resultado.arquivo
                print(f"[{posicao}/{len(fila)}] {chave} baixado -> {resultado.arquivo}")
            elif resultado.status == STATUS_NAO_ENCONTRADO:
                print(f"[{posicao}/{len(fila)}] {chave} não encontrado no portal")
            else:
                # Screenshot antes de fechar o toast, para ele aparecer na imagem.
                salvar_screenshot(sessao.page, chave)
                fechar_notificacoes(sessao.page)
                registrar_erro(chave, resultado.erro)
                erros_rodada += 1
                print(f"[{posicao}/{len(fila)}] {chave} AVISO - ERRO: {truncar(resultado.erro, 160)}")
                print("    seguindo para a próxima; será reprocessada na próxima execução.")
    except KeyboardInterrupt:
        interrompido = True
        print("\nInterrompido pelo usuário. Gravando o controle...")
    finally:
        gravar_controle(args.controle, controle)
        sessao.fechar()

    imprimir_resumo(controle, "Resumo parcial" if interrompido else "Resumo")
    if erros_rodada:
        print(
            f"\nAVISO: {erros_rodada} nota(s) com erro nesta rodada. Detalhes em {LOG_ERROS}"
            f" e prints em {DIR_SCREENSHOTS}/.\nRode o robô de novo para reprocessá-las."
        )
    return 130 if interrompido else 0


# ---------------------------------------------------------------------------
# 9. Comandos restantes e main()
# ---------------------------------------------------------------------------


def comando_importar(args: argparse.Namespace) -> int:
    resultado = importar_chaves_planilha(args.xlsm, args.aba)
    gravar_chaves(args.saida, resultado.chaves)

    print(f"Lidas       : {resultado.total_lido}")
    print(f"Aceitas     : {len(resultado.chaves)}")
    print(f"Inválidas   : {len(resultado.invalidas)}")
    print(f"Duplicadas  : {resultado.duplicadas}")
    for numero, valor in resultado.invalidas[:20]:
        print(f"  linha {numero}: {valor!r}")
    if len(resultado.invalidas) > 20:
        print(f"  ... e mais {len(resultado.invalidas) - 20} linha(s) inválida(s).")
    print(f"\n{args.saida} gravado com {len(resultado.chaves)} chave(s).")
    return 0


def comando_conferir(args: argparse.Namespace) -> int:
    controle = carregar_controle(args.controle)

    # Na primeira execução o controle ainda não existe; semear com chaves.csv é
    # o que torna a reconciliação da pasta histórica útil antes do primeiro
    # `baixar`. Se o arquivo de chaves não existir, apenas segue.
    if Path(args.chaves).exists():
        novas = sincronizar_controle(controle, ler_chaves(args.chaves))
        if novas:
            print(f"{novas} chave(s) de {args.chaves} inserida(s) como pendente.")

    em_disco, sem_chave = varrer_destino(args.destino)

    com_arquivo = 0
    for chave, arquivo in em_disco.items():
        reg = controle.get(chave)
        if reg is None:
            continue
        com_arquivo += 1
        if reg.status != STATUS_BAIXADO or reg.arquivo != arquivo.name:
            reg.marcar(STATUS_BAIXADO, arquivo=arquivo.name)

    sem_arquivo = len(controle) - com_arquivo
    orfaos = [c for c in em_disco if c not in controle]
    gravar_controle(args.controle, controle)

    print(f"Chaves com arquivo em disco : {com_arquivo}")
    print(f"Chaves sem arquivo          : {sem_arquivo}")
    print(f"Arquivos órfãos             : {len(orfaos) + len(sem_chave)}"
          f" ({len(orfaos)} com chave fora do controle, {len(sem_chave)} sem chave no nome)")
    imprimir_resumo(controle)
    return 0


def comando_status(args: argparse.Namespace) -> int:
    controle = carregar_controle(args.controle)
    if not controle:
        print(f"Controle vazio ou inexistente: {args.controle}")
        return 0
    imprimir_resumo(controle, f"Status ({args.controle})")
    erros = [r for r in controle.values() if r.status == STATUS_ERRO]
    if erros:
        print("\nÚltimos erros:")
        for reg in erros[:10]:
            print(f"  {reg.chave}  {truncar(reg.erro, 100)}")
        if len(erros) > 10:
            print(f"  ... e mais {len(erros) - 10}.")
    return 0


def comando_inspecionar(args: argparse.Namespace) -> int:
    cfg = carregar_config()
    timeout_ms = int(args.timeout * 1000)
    sessao = Sessao(cfg, Path(DESTINO_PADRAO), False, timeout_ms, args.perfil, args.espera_login)
    try:
        sessao.abrir()
        if "emissao/notas" not in (sessao.page.url or ""):
            sessao.page.goto(URL_NOTAS, wait_until="domcontentloaded")
        print("Playwright Inspector aberto. Use 'Pick locator' para achar os seletores.")
        print("Feche o Inspector (ou clique em Resume) para encerrar.")
        sessao.page.pause()
    finally:
        sessao.fechar()
    return 0


def construir_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="download_nfce.py",
        description="Baixa PDFs de NFC-e do portal Tec Fiscal / Linx.",
    )
    sub = parser.add_subparsers(dest="comando", required=True)

    p_imp = sub.add_parser("importar", help="Gera chaves.csv a partir de um .xlsm")
    p_imp.add_argument("--xlsm", required=True, help="Planilha de origem")
    p_imp.add_argument("--aba", default="Base", help="Aba com as chaves (padrão: Base)")
    p_imp.add_argument("--saida", default=CHAVES_PADRAO, help=f"Padrão: {CHAVES_PADRAO}")
    p_imp.set_defaults(func=comando_importar)

    p_bai = sub.add_parser("baixar", help="Percorre a fila e baixa os PDFs")
    p_bai.add_argument("--chaves", default=CHAVES_PADRAO, help=f"Padrão: {CHAVES_PADRAO}")
    p_bai.add_argument("--controle", default=CONTROLE_PADRAO, help=f"Padrão: {CONTROLE_PADRAO}")
    p_bai.add_argument("--destino", default=DESTINO_PADRAO, help=f"Padrão: {DESTINO_PADRAO}")
    p_bai.add_argument("--limite", type=int, default=0, help="Máximo de chaves nesta rodada")
    p_bai.add_argument("--lote", type=int, default=LOTE_PADRAO, help="Recicla a sessão a cada N chaves")
    p_bai.add_argument(
        "--tentativas",
        type=int,
        default=TENTATIVAS_PADRAO,
        help=f"Tentativas por chave; a cada {TENTATIVAS_ANTES_RECARREGAR} erros seguidos "
             "a página é recarregada (padrão: %(default)s)",
    )
    p_bai.add_argument("--timeout", type=float, default=TIMEOUT_PADRAO, help="Timeout em segundos")
    p_bai.add_argument("--retentar", action="store_true", help="Inclui também as chaves nao_encontrado")
    p_bai.add_argument("--headless", action="store_true", help="Sem janela (exige login automático)")
    p_bai.add_argument(
        "--perfil",
        default=None,
        metavar="DIR",
        help="Diretório de perfil do Chromium. Preserva a sessão entre reciclagens "
             "e execuções — necessário na prática, porque o login tem captcha.",
    )
    p_bai.add_argument(
        "--formato", choices=("danfe", "xml"), default="danfe",
        help="danfe = PDF (padrão); xml = arquivo fiscal. Use um --destino por "
             "formato: a checagem de idempotência é por chave, não por extensão.",
    )
    p_bai.add_argument(
        "--espera-login", type=int, default=ESPERA_LOGIN_MANUAL, metavar="S",
        help=f"Segundos concedidos ao login manual (padrão: {ESPERA_LOGIN_MANUAL})",
    )
    p_bai.set_defaults(func=comando_baixar)

    p_con = sub.add_parser("conferir", help="Reconcilia o controle com os arquivos em disco")
    p_con.add_argument("--chaves", default=CHAVES_PADRAO, help=f"Semeia o controle. Padrão: {CHAVES_PADRAO}")
    p_con.add_argument("--controle", default=CONTROLE_PADRAO, help=f"Padrão: {CONTROLE_PADRAO}")
    p_con.add_argument("--destino", default=DESTINO_PADRAO, help=f"Padrão: {DESTINO_PADRAO}")
    p_con.set_defaults(func=comando_conferir)

    p_sta = sub.add_parser("status", help="Contagem por status")
    p_sta.add_argument("--controle", default=CONTROLE_PADRAO, help=f"Padrão: {CONTROLE_PADRAO}")
    p_sta.set_defaults(func=comando_status)

    p_ins = sub.add_parser("inspecionar", help="Abre o navegador logado com o Playwright Inspector")
    p_ins.add_argument("--timeout", type=float, default=TIMEOUT_PADRAO, help="Timeout em segundos")
    p_ins.add_argument("--perfil", default=None, metavar="DIR", help="Diretório de perfil do Chromium")
    p_ins.add_argument(
        "--espera-login", type=int, default=ESPERA_LOGIN_MANUAL, metavar="S",
        help=f"Segundos concedidos ao login manual (padrão: {ESPERA_LOGIN_MANUAL})",
    )
    p_ins.set_defaults(func=comando_inspecionar)

    return parser


def main(argv: list[str] | None = None) -> int:
    args = construir_parser().parse_args(argv)
    try:
        return args.func(args)
    except KeyboardInterrupt:
        print("\nInterrompido.")
        return 130


if __name__ == "__main__":
    sys.exit(main())
